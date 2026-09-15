from pathlib import Path

from openpyxl import load_workbook

from app.services.parsing.base import ParseResult, ParsedBlock, Parser, register


@register(".xlsx")
class XlsxParser(Parser):
    def parse(self, path: Path) -> ParseResult:
        result = ParseResult()
        wb = load_workbook(str(path), data_only=True, read_only=True)
        try:
            for idx, ws in enumerate(wb.worksheets, start=1):
                lines = []
                for row in ws.iter_rows(values_only=True):
                    if row is None:
                        continue
                    cells = ["" if v is None else str(v) for v in row]
                    if any(c.strip() for c in cells):
                        lines.append(" | ".join(cells))
                if lines:
                    result.blocks.append(
                        ParsedBlock(content="\n".join(lines), page_no=idx, is_table=True)
                    )
                    result.page_count = idx
        finally:
            wb.close()
        return result
