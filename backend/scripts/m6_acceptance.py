"""M6 无头验收脚本(真栈:http://127.0.0.1:8001 + worker + 智谱 key)。

用法(backend 目录,项目 venv):
    .venv\\Scripts\\python scripts\\m6_acceptance.py

覆盖 M6 四特性:多跳零命中兜底 / 复合问题质量 / eval_generation LLM-judge /
审计留存清理(admin purge 端点)。ZHIPU_API_KEY 未配置时第 3/4 项 SKIP
(决策逻辑单元测试已覆盖)。
"""
import asyncio
import io
import json
import subprocess
import sys
import time
from pathlib import Path

import httpx
from docx import Document as DocxDocument

BASE = "http://127.0.0.1:8001/api"
TIMEOUT = httpx.Timeout(300.0)
RESULTS = {"pass": [], "fail": [], "skip": []}

# 审计旧行直插用的开发库 DSN(与 app 配置一致,绕 ORM 便于设置 created_at)
DSN = "postgresql://airag:airag_dev_password@localhost:5432/airag"

KB_NAME = "M6验收库"
FACT1 = "猎户座反应堆功率五十兆瓦"
FACT2 = "北极星货运飞船载重三十吨"
ZERO_HIT_Q = "企鹅的巡游路线和考拉的食谱是什么"  # 库内零命中词
COMPOSITE_Q = "猎户座反应堆的功率和北极星飞船的载重分别是多少"
DOCX_MIME = ("application/vnd.openxmlformats-officedocument"
             ".wordprocessingml.document")


def check(name, cond, detail=""):
    if cond:
        RESULTS["pass"].append(name)
        print(f"PASS {name}")
    else:
        RESULTS["fail"].append(name)
        print(f"FAIL {name}  {detail}")


def skip(name, why):
    RESULTS["skip"].append(name)
    print(f"SKIP {name}  ({why})")


def promote_roles(usernames: dict[str, str]) -> None:
    async def _run():
        from sqlalchemy import text

        from app.db.session import SessionLocal

        async with SessionLocal() as s:
            for username, role in usernames.items():
                await s.execute(
                    text("UPDATE users SET role = :r WHERE username = :u"),
                    {"r": role, "u": username},
                )
            await s.commit()

    asyncio.run(_run())


def seed_old_audit_rows(n: int = 2) -> None:
    """asyncpg 直插 n 条 200 天前的审计行(超过 180 天保留期,应被 purge 删)。"""

    async def _run():
        import asyncpg

        conn = await asyncpg.connect(DSN)
        try:
            await conn.execute(
                "INSERT INTO audit_logs (username, action, target, created_at) "
                "SELECT 'm6_seed', 'login_success', 'seed-old', "
                "now() - interval '200 days' FROM generate_series(1, $1)",
                n,
            )
        finally:
            await conn.close()

    asyncio.run(_run())


def register_and_login(client: httpx.Client, username: str) -> dict:
    client.post(f"{BASE}/auth/register",
                json={"username": username, "password": "secret123"})
    r = client.post(f"{BASE}/auth/login",
                    json={"username": username, "password": "secret123"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def make_docx_bytes(title: str, paragraphs: list[str]) -> bytes:
    """python-docx 造真 docx(内存字节,不落盘)。"""
    doc = DocxDocument()
    doc.add_heading(title, 0)
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def upload_docx(client: httpx.Client, headers: dict, kb_id: int,
                filename: str, content: bytes) -> dict:
    r = client.post(
        f"{BASE}/kbs/{kb_id}/documents",
        files={"file": (filename, content, DOCX_MIME)},
        data={"ocr": "auto"},
        headers=headers,
    )
    return {"status": r.status_code, "body": r.json(), "text": r.text[:200]}


def wait_doc_status(client: httpx.Client, headers: dict, doc_id: int,
                    terminal=("done", "failed"), timeout_s=120) -> dict:
    deadline = time.time() + timeout_s
    while time.time() < deadline:
        r = client.get(f"{BASE}/documents/{doc_id}", headers=headers)
        doc = r.json()
        if doc["status"] in terminal:
            return doc
        time.sleep(2)
    return {"status": "timeout", "error_msg": "poll timeout"}


def sse_ask(client: httpx.Client, headers: dict, kb_ids: list[int],
            question: str) -> dict:
    """POST /chat/ask,消费 SSE,返回 {tokens, citations, done, error}。"""
    out = {"tokens": "", "citations": None, "done": None, "error": None}
    payload = {"kb_ids": kb_ids, "question": question}
    with client.stream("POST", f"{BASE}/chat/ask", json=payload,
                       headers=headers, timeout=TIMEOUT) as resp:
        out["status"] = resp.status_code
        buf = b""
        for chunk in resp.iter_bytes():
            buf += chunk
            while b"\n\n" in buf:
                frame, buf = buf.split(b"\n\n", 1)
                line = frame.decode("utf-8").strip()
                if not line.startswith("data:"):
                    continue
                evt = json.loads(line[5:])
                if evt["type"] == "token":
                    out["tokens"] += str(evt["data"])
                elif evt["type"] == "citations":
                    out["citations"] = evt["data"]
                elif evt["type"] == "done":
                    out["done"] = evt["data"]
                elif evt["type"] == "error":
                    out["error"] = evt["data"]
                    return out
    return out


def final_answer(res: dict) -> str:
    """done 帧携带权威终稿;流式 token 作对账兜底。"""
    return ((res.get("done") or {}).get("answer") or res["tokens"]) or ""


def main():
    from app.core.config import settings

    client = httpx.Client(timeout=60.0)
    has_llm = bool(settings.ZHIPU_API_KEY)

    # ---- 1. 基线 ----
    r = client.get(f"{BASE}/health")
    check("health 200", r.status_code == 200)
    client.post(f"{BASE}/auth/register",
                json={"username": "m6_user", "password": "secret123"})
    client.post(f"{BASE}/auth/register",
                json={"username": "m6_owner", "password": "secret123"})
    promote_roles({"m6_user": "admin", "m6_owner": "editor"})
    admin = register_and_login(client, "m6_user")
    owner = register_and_login(client, "m6_owner")
    me = client.get(f"{BASE}/auth/me", headers=owner).json()
    check("m6_owner promoted editor", me["role"] == "editor", str(me))

    # ---- 2. 零命中多跳路径(确定性,验证兜底不炸)----
    r = client.post(f"{BASE}/kbs", json={"name": KB_NAME}, headers=owner)
    kb_id = r.json()["id"]
    check("kb created", r.status_code == 201 and kb_id > 0, r.text[:200])

    up = upload_docx(client, owner, kb_id, "m6猎户座.docx",
                     make_docx_bytes("猎户座反应堆简介", [FACT1]))
    doc1_id = up["body"].get("id")
    doc1 = wait_doc_status(client, owner, doc1_id)
    check("docx#1 upload & processed done",
          up["status"] == 201 and doc1["status"] == "done"
          and doc1["chunk_count"] > 0, f"{up} {doc1}")

    zero = sse_ask(client, owner, [kb_id], ZERO_HIT_Q)
    check("zero-hit sse done frame, no error",
          zero["done"] is not None and zero["error"] is None, str(zero)[:300])
    zero_ans = final_answer(zero)
    check("zero-hit answer mentions 未找到", "未找到" in zero_ans,
          zero_ans[:200])

    # ---- 3. 复合问题质量(真智谱 LLM,失败重试 1 次)----
    if not has_llm:
        skip("composite ask contains 五十 and 三十", "ZHIPU_API_KEY 未配置")
        doc2_id = None
    else:
        up2 = upload_docx(client, owner, kb_id, "m6北极星.docx",
                          make_docx_bytes("北极星货运飞船简介", [FACT2]))
        doc2_id = up2["body"].get("id")
        doc2 = wait_doc_status(client, owner, doc2_id)
        check("docx#2 upload & processed done",
              up2["status"] == 201 and doc2["status"] == "done"
              and doc2["chunk_count"] > 0, f"{up2} {doc2}")
        comp_ok, comp_ans = False, ""
        for _attempt in range(2):
            comp = sse_ask(client, owner, [kb_id], COMPOSITE_Q)
            comp_ans = final_answer(comp)
            if comp["done"] and "五十" in comp_ans and "三十" in comp_ans:
                comp_ok = True
                break
        check("composite ask contains 五十 and 三十", comp_ok, comp_ans[:200])

    # ---- 4. eval_generation CLI(真智谱)----
    if not has_llm:
        skip("eval cli 3 items judged, faithfulness avg >= 0.8",
             "ZHIPU_API_KEY 未配置")
    else:
        eval_path = Path("eval_sets") / f"{kb_id}.json"
        eval_path.parent.mkdir(exist_ok=True)
        eval_path.write_text(json.dumps({
            "kb_id": kb_id,
            "items": [
                {"question": "猎户座反应堆的功率是多少?",
                 "expect_doc_ids": [doc1_id, doc2_id],
                 "expect_keywords": ["五十"]},
                {"question": "北极星货运飞船的载重是多少?",
                 "expect_doc_ids": [doc1_id, doc2_id],
                 "expect_keywords": ["三十"]},
                {"question": COMPOSITE_Q,
                 "expect_doc_ids": [doc1_id, doc2_id],
                 "expect_keywords": ["五十", "三十"]},
            ],
        }, ensure_ascii=False), encoding="utf-8")
        proc = subprocess.run(
            [sys.executable, "-m", "scripts.eval_generation",
             "--kb", str(kb_id), "--json"],
            capture_output=True, text=True, timeout=600,
        )
        try:
            eval_out = json.loads(proc.stdout)
            faith = [i["faithfulness"]["score"] for i in eval_out]
            relev = [i["relevancy"]["score"] for i in eval_out]
            check("eval cli 3 items judged, faithfulness avg >= 0.8",
                  len(eval_out) == 3
                  and all(s is not None for s in faith + relev)
                  and sum(faith) / len(faith) >= 0.8,
                  proc.stdout[:300])
        except Exception as exc:
            check("eval cli 3 items judged, faithfulness avg >= 0.8", False,
                  f"{exc}; stdout={proc.stdout[:200]} stderr={proc.stderr[:200]}")

    # ---- 5. 审计留存清理 ----
    seed_old_audit_rows(2)
    r = client.post(f"{BASE}/admin/audit-logs/purge", headers=admin)
    body = r.json() if r.status_code == 200 else {}
    check("purge deleted >= 2 old rows, retention_days == 180",
          r.status_code == 200 and body.get("deleted", 0) >= 2
          and body.get("retention_days") == 180, f"{r.status_code} {r.text[:200]}")
    r = client.get(f"{BASE}/admin/audit-logs", params={"action": "audit_purge"},
                   headers=admin)
    check("audit_purge action logged",
          r.status_code == 200 and r.json().get("total", 0) >= 1, r.text[:200])
    r = client.post(f"{BASE}/admin/audit-logs/purge", headers=owner)
    check("editor (non-admin) purge forbidden 403", r.status_code == 403,
          f"{r.status_code} {r.text[:100]}")

    # ---- 6. 汇总 ----
    print()
    total = sum(len(v) for v in RESULTS.values())
    print(f"M6 ACCEPTANCE: {len(RESULTS['pass'])}/{total} PASS")
    if RESULTS["skip"]:
        print("SKIP:", *RESULTS["skip"], sep="\n  - ")
    if RESULTS["fail"]:
        print("FAILED:", *RESULTS["fail"], sep="\n  - ")
        sys.exit(1)
    if not has_llm:
        print("提示:配置 ZHIPU_API_KEY 后重跑,可覆盖 SKIP 的复合问题质量与"
              " eval_generation 两项。")


if __name__ == "__main__":
    main()
