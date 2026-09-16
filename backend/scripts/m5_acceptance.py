"""M5 无头验收脚本(真栈:http://127.0.0.1:8001 + worker + 智谱 key)。

用法(backend 目录,项目 venv):
    .venv\\Scripts\\python scripts\\m5_acceptance.py

MINERU_API_TOKEN 未配置时 OCR 真调项自动 SKIP(mock 单测已覆盖决策逻辑)。
"""
import asyncio
import io
import json
import subprocess
import sys
import time
from pathlib import Path

import httpx
import pymupdf as fitz
from docx import Document as DocxDocument

BASE = "http://127.0.0.1:8001/api"
TIMEOUT = httpx.Timeout(300.0)
RESULTS = {"pass": [], "fail": [], "skip": []}

KB_NAME = "M5验收库"
FACTS = {
    "budget": "星际探索项目预算总额为三千万元",
    "owner": "项目负责人是张三丰",
    "period": "项目周期为2026至2028年",
}


def check(name, cond, detail=""):
    if cond:
        RESULTS["pass"].append(name)
        print(f"PASS {name}")
    else:
        RESULTS["fail"].append(name)
        print(f"FAIL {name}  {detail}")


def skip(name, why):
    RESULTS["skip"].append(name)
    print(f"SKIP {name}  ({why})")


def promote_roles(usernames: dict[str, str]) -> None:
    async def _run():
        from sqlalchemy import text

        from app.db.session import SessionLocal

        async with SessionLocal() as s:
            for username, role in usernames.items():
                await s.execute(
                    text("UPDATE users SET role = :r WHERE username = :u"),
                    {"r": role, "u": username},
                )
            await s.commit()

    asyncio.run(_run())


def register_and_login(client: httpx.Client, username: str) -> dict:
    client.post(f"{BASE}/auth/register",
                json={"username": username, "password": "secret123"})
    r = client.post(f"{BASE}/auth/login",
                    json={"username": username, "password": "secret123"})
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def make_fact_docx(path: Path) -> None:
    doc = DocxDocument()
    doc.add_heading("星际探索项目简介", 0)
    doc.add_paragraph(FACTS["budget"])
    doc.add_paragraph(FACTS["owner"])
    doc.add_paragraph(FACTS["period"])
    doc.save(str(path))


def make_thin_pdf(path: Path) -> None:
    """零文本层 PDF(空页)——ocr=off 时应解析失败。"""
    doc = fitz.open()
    doc.new_page(width=400, height=400)
    doc.save(str(path))


def poll_status(client: httpx.Client, headers: dict, doc_id: int,
                terminal=("done", "failed"), timeout_s=360) -> dict:
    deadline = time.time() + timeout_s
    while time.time() < deadline:
        r = client.get(f"{BASE}/documents/{doc_id}", headers=headers)
        doc = r.json()
        if doc["status"] in terminal:
            return doc
        time.sleep(2)
    return {"status": "timeout", "error_msg": "poll timeout"}


def sse_ask(client: httpx.Client, headers: dict, kb_ids: list[int],
            question: str, conversation_id: int | None = None) -> dict:
    """POST /chat/ask,消费 SSE,返回 {tokens, citations, done, error}。"""
    out = {"tokens": "", "citations": None, "done": None, "error": None}
    payload = {"kb_ids": kb_ids, "question": question}
    if conversation_id is not None:
        payload["conversation_id"] = conversation_id
    with client.stream("POST", f"{BASE}/chat/ask", json=payload,
                       headers=headers, timeout=TIMEOUT) as resp:
        out["status"] = resp.status_code
        buf = b""
        for chunk in resp.iter_bytes():
            buf += chunk
            while b"\n\n" in buf:
                frame, buf = buf.split(b"\n\n", 1)
                line = frame.decode("utf-8").strip()
                if not line.startswith("data:"):
                    continue
                evt = json.loads(line[5:])
                if evt["type"] == "token":
                    out["tokens"] += str(evt["data"])
                elif evt["type"] == "citations":
                    out["citations"] = evt["data"]
                elif evt["type"] == "done":
                    out["done"] = evt["data"]
                elif evt["type"] == "error":
                    out["error"] = evt["data"]
                    return out
    return out


def main():
    from app.core.config import settings

    client = httpx.Client(timeout=60.0)

    # ---- 1. 基线 ----
    r = client.get(f"{BASE}/health")
    check("health 200", r.status_code == 200)
    client.post(f"{BASE}/auth/register",
                json={"username": "m5_user", "password": "secret123"})
    client.post(f"{BASE}/auth/register",
                json={"username": "m5_owner", "password": "secret123"})
    promote_roles({"m5_user": "admin", "m5_owner": "editor"})
    admin = register_and_login(client, "m5_user")
    owner = register_and_login(client, "m5_owner")
    me = client.get(f"{BASE}/auth/me", headers=owner).json()
    check("m5_owner promoted editor", me["role"] == "editor", str(me))

    # ---- 2. 建库 + 事实 docx 上传到 done(真嵌入)----
    r = client.post(f"{BASE}/kbs", json={"name": KB_NAME}, headers=owner)
    kb_id = r.json()["id"]
    check("kb created", r.status_code == 201 and kb_id > 0)

    fact_docx = Path("storage/m5_fact.docx")
    fact_docx.parent.mkdir(parents=True, exist_ok=True)
    make_fact_docx(fact_docx)
    with open(fact_docx, "rb") as f:
        r = client.post(
            f"{BASE}/kbs/{kb_id}/documents",
            files={"file": ("m5事实文档.docx", f,
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"ocr": "auto"},
            headers=owner,
        )
    doc_id = r.json()["id"]
    check("docx upload 201 with ocr fields",
          r.status_code == 201 and r.json()["ocr_mode"] == "auto"
          and "ocr_used" in r.json(), r.text[:200])
    doc = poll_status(client, owner, doc_id)
    check("docx processed done", doc["status"] == "done"
          and doc["chunk_count"] > 0, str(doc))
    r = client.get(f"{BASE}/documents/{doc_id}/chunks", headers=owner)
    blob = json.dumps([i["content_preview"] for i in r.json()["items"]],
                      ensure_ascii=False)
    check("chunks contain budget fact", "三千" in blob)

    # ---- 3. ocr=off 对照:零文本层 PDF → failed ----
    thin_pdf = Path("storage/m5_thin.pdf")
    make_thin_pdf(thin_pdf)
    with open(thin_pdf, "rb") as f:
        r = client.post(
            f"{BASE}/kbs/{kb_id}/documents",
            files={"file": ("m5空白页.pdf", f, "application/pdf")},
            data={"ocr": "off"},
            headers=owner,
        )
    thin_id = r.json()["id"]
    thin_doc = poll_status(client, owner, thin_id)
    check("thin pdf with ocr=off fails as expected",
          thin_doc["status"] == "failed", str(thin_doc))

    # ---- 4. OCR 真调(需 MINERU_API_TOKEN)----
    if not settings.MINERU_API_TOKEN:
        skip("scanned pdf via mineru (auto)", "MINERU_API_TOKEN 未配置")
        skip("image png via mineru (auto)", "MINERU_API_TOKEN 未配置")
    else:
        # 图片型 PDF:文字渲染成位图再插入(无文本层),用 pymupdf 免 PIL 依赖
        img_doc = fitz.open()
        img_page = img_doc.new_page(width=612, height=792)
        y = 80
        for line in (FACTS["budget"], FACTS["owner"], FACTS["period"]):
            img_page.insert_text((50, y), line, fontsize=14, fontname="china-s")
            y += 40
        pix = img_page.get_pixmap(matrix=fitz.Matrix(2, 2))
        scan_pdf = Path("storage/m5_scan.pdf")
        pdf = fitz.open()
        page = pdf.new_page(width=612, height=792)
        page.insert_image(fitz.Rect(0, 0, 612, 792), pixmap=pix)
        pdf.save(str(scan_pdf))
        with open(scan_pdf, "rb") as f:
            r = client.post(
                f"{BASE}/kbs/{kb_id}/documents",
                files={"file": ("m5扫描件.pdf", f, "application/pdf")},
                data={"ocr": "auto"},
                headers=owner,
            )
        scan_id = r.json()["id"]
        scan_doc = poll_status(client, owner, scan_id, timeout_s=360)
        check("scanned pdf ocr done + ocr_used",
              scan_doc["status"] == "done" and scan_doc["ocr_used"] is True,
              str(scan_doc))
        r = client.get(f"{BASE}/documents/{scan_id}/chunks", headers=owner)
        blob = json.dumps([i["content_preview"] for i in r.json()["items"]],
                          ensure_ascii=False)
        check("ocr chunks contain fact", "三千" in blob or "扫描" in blob)

    # ---- 5. agentic 多轮问答(真 LLM,rewrite 默认开)----
    first = sse_ask(client, owner, [kb_id], "星际探索项目的预算总额是多少?")
    check("ask#1 sse token->citations->done",
          first["done"] is not None and first["error"] is None
          and first["citations"], str(first)[:300])
    conv_id = (first["done"] or {}).get("conversation_id")
    check("ask#1 answer mentions 三千", "三千" in first["tokens"],
          first["tokens"][:200])

    second_ok = False
    for attempt in range(2):
        second = sse_ask(client, owner, [kb_id], "它的负责人是谁?",
                         conversation_id=conv_id)
        if second["done"] and "张三丰" in second["tokens"]:
            second_ok = True
            break
    check("ask#2 coreference rewritten and answered 张三丰",
          second_ok, second["tokens"][:200])

    # ---- 6. 审计矩阵 ----
    client.post(f"{BASE}/auth/login",
                json={"username": "m5_owner", "password": "wrong"})
    client.post(f"{BASE}/auth/register",
                json={"username": "m5_grantee", "password": "secret123"})
    client.put(f"{BASE}/kbs/{kb_id}/permissions",
               json={"username": "m5_grantee", "perm": "viewer"}, headers=owner)
    client.delete(f"{BASE}/kbs/{kb_id}/permissions",
                  params={"username": "m5_grantee"}, headers=owner)
    conv_tmp = client.post(f"{BASE}/chat/conversations",
                           json={"kb_ids": [kb_id], "name": "审计用"}, headers=owner)
    client.delete(f"{BASE}/chat/conversations/{conv_tmp.json()['id']}",
                  headers=owner)
    grantee = client.get(f"{BASE}/auth/me", headers=register_and_login(
        client, "m5_grantee")).json()
    client.patch(f"{BASE}/admin/users/{grantee['id']}",
                 json={"role": "editor"}, headers=admin)

    r = client.get(f"{BASE}/admin/audit-logs", params={"page_size": 100},
                   headers=admin)
    check("audit api 200", r.status_code == 200)
    actions = {i["action"] for i in r.json()["items"]}
    need = {"login_success", "login_fail", "register", "kb_create",
            "kb_grant", "kb_revoke", "doc_upload", "conv_delete",
            "user_admin_update", "ask"}
    check("audit covers all actions", need <= actions,
          f"missing: {need - actions}")
    r = client.get(f"{BASE}/admin/audit-logs", params={"username": "m5_owner"},
                   headers=admin)
    check("audit filter by username", r.json()["total"] > 0
          and all(i["username"] == "m5_owner" for i in r.json()["items"]))
    r = client.get(f"{BASE}/admin/audit-logs", params={"page": 1, "page_size": 1},
                   headers=admin)
    check("audit pagination", len(r.json()["items"]) == 1)
    r = client.get(f"{BASE}/admin/audit-logs", headers=owner)
    check("audit admin-only 403", r.status_code == 403)

    # ---- 7. 会话导出 ----
    r = client.get(f"{BASE}/chat/conversations/{conv_id}/export", headers=owner)
    body = r.text
    check("export markdown with appendix",
          r.status_code == 200 and "预算总额" in body
          and "引用附录" in body and "m5事实文档.docx" in body, body[:200])
    r = client.get(f"{BASE}/chat/conversations/{conv_id}/export", headers=admin)
    check("export owner-only 404", r.status_code == 404)

    # ---- 8. 评估集 CLI ----
    eval_path = Path("eval_sets") / f"{kb_id}.json"
    eval_path.parent.mkdir(exist_ok=True)
    eval_path.write_text(json.dumps({
        "kb_id": kb_id,
        "items": [
            {"question": "星际探索项目的预算总额是多少?",
             "expect_doc_ids": [doc_id], "expect_keywords": ["三千"]},
            {"question": "星际探索项目的负责人是谁?",
             "expect_doc_ids": [doc_id], "expect_keywords": ["张三丰"]},
            {"question": "项目周期是什么时候?",
             "expect_doc_ids": [doc_id], "expect_keywords": ["2026"]},
        ],
    }, ensure_ascii=False), encoding="utf-8")
    proc = subprocess.run(
        [sys.executable, "-m", "scripts.eval_retrieval", "--kb", str(kb_id),
         "--json"],
        capture_output=True, text=True, timeout=120,
    )
    try:
        eval_out = json.loads(proc.stdout)
        hit_all = all(item["hit_at_k"] for item in eval_out)
        mrr_all = all(item["mrr"] == 1.0 for item in eval_out)
        check("eval cli 3 items all hit mrr=1",
              len(eval_out) == 3 and hit_all and mrr_all, proc.stdout[:300])
    except Exception as exc:
        check("eval cli 3 items all hit mrr=1", False,
              f"{exc}; stdout={proc.stdout[:200]} stderr={proc.stderr[:200]}")

    # ---- 9. Minor 抽查 ----
    r = client.get(f"{BASE}/kbs", headers=owner)
    row = next((k for k in r.json() if k["id"] == kb_id), {})
    check("kb doc_count >= 1", row.get("doc_count", 0) >= 1, str(row))

    # ---- 汇总 ----
    print()
    print(f"M5 ACCEPTANCE: {len(RESULTS['pass'])} PASS / "
          f"{len(RESULTS['fail'])} FAIL / {len(RESULTS['skip'])} SKIP")
    if RESULTS["fail"]:
        print("FAILED:", *RESULTS["fail"], sep="\n  - ")
        sys.exit(1)


if __name__ == "__main__":
    main()
