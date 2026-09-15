import io
from pathlib import Path

import pymupdf as fitz
import pytest
from docx import Document as Docx
from openpyxl import Workbook

from app.services.parsing.base import get_parser


def _make_pdf(tmp_path: Path) -> Path:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "AIRag pdf parser test page one")
    page2 = doc.new_page()
    page2.insert_text((72, 72), "page two content")
    p = tmp_path / "t.pdf"
    doc.save(str(p))
    doc.close()
    return p


def _make_docx(tmp_path: Path) -> Path:
    d = Docx()
    d.add_paragraph("第一段介绍")
    d.add_paragraph("")
    d.add_paragraph("第二段内容")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "90"
    p = tmp_path / "t.docx"
    d.save(str(p))
    return p


def _make_xlsx(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Score"])
    ws.append(["Alice", "90"])
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["City", "Code"])
    ws2.append(["BJ", "010"])
    p = tmp_path / "t.xlsx"
    wb.save(str(p))
    return p


def test_pdf_parses_pages(tmp_path):
    result = get_parser(".pdf").parse(_make_pdf(tmp_path))
    assert result.page_count == 2
    assert result.blocks[0].page_no == 1
    assert "page one" in result.blocks[0].content
    assert result.blocks[1].content == "page two content"


def test_docx_keeps_order_and_table(tmp_path):
    result = get_parser(".docx").parse(_make_docx(tmp_path))
    texts = [b.content for b in result.blocks]
    assert "第一段介绍" in texts
    tables = [b for b in result.blocks if b.is_table]
    assert len(tables) == 1
    assert "Name | Score" in tables[0].content
    assert "Alice | 90" in tables[0].content
    assert tables[0].page_no is None


def test_xlsx_sheet_as_whole_block(tmp_path):
    result = get_parser(".xlsx").parse(_make_xlsx(tmp_path))
    assert result.page_count == 2
    assert result.blocks[0].is_table is True
    assert "Name | Score" in result.blocks[0].content
    assert result.blocks[1].page_no == 2
    assert "BJ | 010" in result.blocks[1].content


def test_unknown_ext_raises():
    with pytest.raises(KeyError):
        get_parser(".txt")
